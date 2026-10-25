"""SnapBill by Silsilara — Copyright (c) 2026 Risyad Abdala Ramadhan. All rights reserved.

Membuat 2 template Excel untuk impor data ke SnapBill:
  templates/snapbill-template-data-master.xlsx  (Profil Usaha, Klien, Item & Harga)
  templates/snapbill-template-lengkap.xlsx      (+ Dokumen, Rincian)

Jalankan:  python tools/make_excel_templates.py
Butuh:     pip install openpyxl

PENTING: nama sheet & judul kolom harus sama dengan yang dibaca js/excel-import.js.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'templates')

BRAND = '0B7A5F'
HEAD_FILL = PatternFill('solid', fgColor=BRAND)
HINT_FILL = PatternFill('solid', fgColor='F2F7F5')
EX_HEAD_FILL = PatternFill('solid', fgColor='D9EAD3')
THIN = Side(style='thin', color='C9D6D0')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATA_ROWS = 500  # jumlah baris yang diberi format & dropdown

MODELS = ['Klasik', 'Grup Lokasi', 'Ringkas', 'Tabel Rincian', 'Kuantitas']
RP = '#,##0'
DATE = 'dd/mm/yyyy'

# (judul kolom, keterangan baris 2, lebar, format angka, daftar dropdown)
SHEETS = {
    'Klien': [
        ('Nama Klien', 'Wajib. Mis. Bu Sinta', 26, None, None),
        ('No. WhatsApp', 'Mis. 0812xxxx', 18, '@', None),
        ('Alamat', 'Opsional', 30, None, None),
        ('Biaya Bulanan (Rp)', 'Angka saja, mis. 450000', 20, RP, None),
        ('Lokasi', 'Pisahkan dengan ; mis. Kantor Lt. 1; Kantor Lt. 2', 38, None, None),
        ('Catatan', 'Opsional', 30, None, None),
    ],
    'Item & Harga': [
        ('Nama Item', 'Wajib. Layanan, barang, atau material', 32, None, None),
        ('Satuan', 'Mis. unit, m, kg, jam, phn, krg', 16, None, None),
        ('Harga Satuan (Rp)', 'Angka saja, mis. 150000', 20, RP, None),
    ],
    'Dokumen': [
        ('No. Dokumen', 'Wajib & unik. Mis. INV/2026/10/001', 22, '@', None),
        ('Jenis', 'Invoice / Penawaran', 14, None, ['Invoice', 'Penawaran']),
        ('Nama Klien', 'Wajib. Sama seperti di sheet Klien', 24, None, None),
        ('Tanggal', 'Wajib. dd/mm/yyyy', 14, DATE, None),
        ('Periode (Bulan)', 'Opsional. Tanggal apa saja di bulan tsb', 16, 'mmmm yyyy', None),
        ('Jatuh Tempo', 'Opsional. dd/mm/yyyy', 14, DATE, None),
        ('Model', 'Kosong = Klasik', 16, None, MODELS),
        ('Harga Borongan (Rp)', 'Isi jika 1 harga untuk semua kunjungan', 20, RP, None),
        ('Diskon (%)', 'Opsional', 11, '0.##', None),
        ('Pajak (%)', 'Opsional', 11, '0.##', None),
        ('Status', 'Draft / Terkirim (kosong = Terkirim)', 14, None, ['Draft', 'Terkirim']),
        ('Catatan', 'Tampil di bawah invoice', 30, None, None),
    ],
    'Rincian': [
        ('No. Dokumen', 'Wajib. Sama seperti di sheet Dokumen', 22, '@', None),
        ('Tipe', 'Kunjungan / Tambahan / DP', 13, None, ['Kunjungan', 'Tambahan', 'DP']),
        ('Lokasi / Bagian', 'Opsional. Mis. Kantor Lt. 1, TAHAP 1', 22, None, None),
        ('Rincian', 'Wajib. Nama layanan/item, atau label DP', 30, None, None),
        ('Tanggal', 'Tanggal kunjungan / bayar DP', 14, DATE, None),
        ('Jumlah', 'Angka, mis. 2 (kosong = 1)', 10, '0.##', None),
        ('Satuan', 'Mis. unit, m, kg, phn', 12, None, None),
        ('Harga Satuan (Rp)', 'Harga per satuan / nominal DP', 20, RP, None),
    ],
}

PROFILE_FIELDS = [
    ('Nama usaha', 'Mis. CV MAJU BERSAMA'),
    ('Slogan', 'Mis. Servis & Instalasi AC'),
    ('No. telepon', 'Mis. 0812 xxxx xxxx'),
    ('Email', 'Opsional'),
    ('Alamat', 'Opsional'),
    ('Nama penanda tangan', 'Mis. BUDI SANTOSO'),
    ('Info pembayaran', 'Mis. BCA 1234567890 a.n. Nama Anda'),
    ('Layanan utama', 'Mis. Servis AC'),
]

# Contoh fiktif (hanya ditampilkan di sheet Petunjuk, tidak ikut terimpor).
EX_DOKUMEN = [
    ['INV/2026/10/001', 'Invoice', 'Bu Sinta', '04/10/2026', 'Oktober 2026', '', 'Klasik', 450000, '', '', 'Terkirim', ''],
    ['PNW/2026/10/001', 'Penawaran', 'PT Sumber Rejeki', '10/10/2026', '', '', 'Kuantitas', '', '', '', 'Draft', 'Berlaku 14 hari'],
]
EX_RINCIAN = [
    ['INV/2026/10/001', 'Kunjungan', '', 'Servis AC', '04/10/2026', '', '', ''],
    ['INV/2026/10/001', 'Kunjungan', '', 'Servis AC', '18/10/2026', '', '', ''],
    ['INV/2026/10/001', 'Tambahan', '', 'Isi freon R32', '', 1, 'kg', 150000],
    ['PNW/2026/10/001', 'Kunjungan', 'TAHAP 1', 'AC Split 1 PK', '', 2, 'unit', 3500000],
    ['PNW/2026/10/001', 'Kunjungan', 'TAHAP 1', 'Pipa AC + insulasi', '', 10, 'm', 120000],
    ['PNW/2026/10/001', 'Tambahan', '', 'Jasa pemasangan', '', 2, 'unit', 350000],
    ['PNW/2026/10/001', 'DP', '', 'DP 1', '12/10/2026', '', '', 4000000],
]


def header_row(ws, cols):
    for i, (title, hint, width, fmt, _) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = BOX
        h = ws.cell(row=2, column=i, value=hint)
        h.font = Font(italic=True, color='6B7C75', size=9)
        h.fill = HINT_FILL
        h.alignment = Alignment(vertical='top', wrap_text=True)
        h.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = width
        if fmt:
            for r in range(3, DATA_ROWS + 3):
                ws.cell(row=r, column=i).number_format = fmt
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = 'A3'


def add_dropdowns(ws, cols):
    for i, (title, _, _, _, options) in enumerate(cols, start=1):
        if not options:
            continue
        dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle=title, error='Pilih dari daftar: ' + ', '.join(options))
        col = get_column_letter(i)
        dv.add(f'{col}3:{col}{DATA_ROWS + 2}')
        ws.add_data_validation(dv)


def data_sheet(wb, name):
    ws = wb.create_sheet(name)
    cols = SHEETS[name]
    header_row(ws, cols)
    add_dropdowns(ws, cols)
    return ws


def profile_sheet(wb):
    ws = wb.create_sheet('Profil Usaha')
    header_row(ws, [('Data', 'Jangan ubah kolom ini', 24, None, None), ('Isi', 'Isi kolom ini', 46, None, None)])
    for r, (label, hint) in enumerate(PROFILE_FIELDS, start=3):
        a = ws.cell(row=r, column=1, value=label)
        a.font = Font(bold=True)
        a.border = BOX
        b = ws.cell(row=r, column=2)
        b.border = BOX
        ws.cell(row=r, column=3, value=hint).font = Font(italic=True, color='8A9A93', size=9)
    ws.column_dimensions['C'].width = 36
    ws.cell(row=len(PROFILE_FIELDS) + 4, column=1, value='Logo diatur langsung di aplikasi (Master → Usaha).').font = Font(italic=True, color='6B7C75')


def example_table(ws, row, title, cols, rows):
    ws.cell(row=row, column=1, value=title).font = Font(bold=True, color=BRAND)
    row += 1
    for i, (t, *_rest) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = Font(bold=True, size=9)
        c.fill = EX_HEAD_FILL
        c.border = BOX
    for data in rows:
        row += 1
        for i, v in enumerate(data, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.font = Font(size=9, color='44544E')
            c.border = BOX
            if isinstance(v, (int, float)) and v >= 1000:
                c.number_format = RP
    return row + 2


def guide_sheet(wb, full):
    ws = wb.active
    ws.title = 'Petunjuk'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 20
    for col in 'BCDEFGHIJKL':
        ws.column_dimensions[col].width = 15
    title = 'Template Impor SnapBill — ' + ('Lengkap' if full else 'Data Master')
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=16, color=BRAND)
    lines = [
        'Cara pakai:',
        '1. Isi sheet di bawah mulai BARIS 3 (baris 1 = judul kolom, baris 2 = keterangan). Jangan ubah nama sheet & judul kolom.',
        '2. Simpan file sebagai .xlsx.',
        '3. Buka SnapBill → Pengaturan → Backup & pindah perangkat → Impor file Excel → pilih file ini.',
        '4. Periksa ringkasan & daftar kesalahan, lalu ketuk "Impor sekarang".',
        '',
        'Isi tiap sheet:',
        '• Profil Usaha — kop invoice: nama usaha, slogan, telepon, penanda tangan, info rekening.',
        '• Klien — daftar klien. "Biaya Bulanan" otomatis jadi harga borongan saat membuat tagihan.',
        '• Item & Harga — layanan, barang, dan material beserta satuan & harga default.',
    ]
    if full:
        lines += [
            '• Dokumen — satu baris per invoice / penawaran (kepala dokumen).',
            '• Rincian — isi tiap dokumen, dihubungkan lewat "No. Dokumen":',
            '      Kunjungan = baris layanan/kunjungan (bisa dikelompokkan per "Lokasi / Bagian", mis. TAHAP 1).',
            '      Tambahan  = item tambahan bersatuan, mis. 1 kg freon × Rp150.000.',
            '      DP        = uang muka. Isi "Rincian" dengan label (DP 1) dan "Harga Satuan" dengan nominal.',
            '• Jika "Harga Borongan" di sheet Dokumen terisi, semua baris Kunjungan memakai satu harga itu.',
        ]
    lines += [
        '',
        'Aturan impor: data dengan nama klien / nama item / nomor dokumen yang sama akan DIPERBARUI,',
        'data lain ditambahkan. Data yang sudah ada di aplikasi tidak dihapus.',
    ]
    r = 3
    for line in lines:
        c = ws.cell(row=r, column=1, value=line)
        c.font = Font(bold=line.endswith(':'), color='1F2D28')
        r += 1
    r += 1
    if full:
        ws.cell(row=r, column=1, value='Contoh pengisian (fiktif — JANGAN disalin ke sheet data):').font = Font(bold=True, italic=True, color='B7791F')
        r += 2
        r = example_table(ws, r, 'Sheet Dokumen', SHEETS['Dokumen'], EX_DOKUMEN)
        example_table(ws, r, 'Sheet Rincian', SHEETS['Rincian'], EX_RINCIAN)
    ws.cell(row=60, column=1, value='© 2026 Risyad Abdala Ramadhan — SnapBill by Silsilara').font = Font(size=8, color='8A9A93')


def build(full):
    wb = Workbook()
    guide_sheet(wb, full)
    profile_sheet(wb)
    data_sheet(wb, 'Klien')
    data_sheet(wb, 'Item & Harga')
    if full:
        data_sheet(wb, 'Dokumen')
        data_sheet(wb, 'Rincian')
    wb.properties.creator = 'SnapBill by Silsilara'
    wb.properties.title = 'Template Impor SnapBill'
    name = 'snapbill-template-lengkap.xlsx' if full else 'snapbill-template-data-master.xlsx'
    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, name)
    wb.save(path)
    return path


if __name__ == '__main__':
    for full in (False, True):
        print('dibuat:', build(full))
